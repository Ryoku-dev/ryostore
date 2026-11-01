"""Fixtures for the Utsuri plugin tests.

Run with the engine's own interpreter, so the tests exercise the exact locked
dependency set the plugin uses at runtime:

    ~/.local/share/utsuri/engine/.venv/bin/python -m pytest tests
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

import pytest

sys.dont_write_bytecode = True  # keep bin/ free of __pycache__ (it gets installed)

PLUGIN = Path(__file__).resolve().parent.parent
BIN = PLUGIN / "bin"
ENGINE = Path(os.environ.get("UTSURI_ENGINE", "~/.local/share/utsuri/engine")).expanduser()

os.environ["UTSURI_ENGINE"] = str(ENGINE)
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(ENGINE))

from utsuri import cli as eng  # noqa: E402

from PIL import Image, ImageDraw  # noqa: E402  (after the engine is on sys.path)


def _gps_exif() -> Image.Exif:
    exif = Image.Exif()
    exif[0x010F] = "TestCam"
    gps = exif.get_ifd(0x8825)
    gps[1], gps[2] = "N", (10.0, 46.0, 30.0)
    gps[3], gps[4] = "E", (106.0, 41.0, 0.0)
    return exif


def _draw(size=(640, 420), mode="RGB") -> Image.Image:
    im = Image.new(mode, size, (250, 250, 250))
    d = ImageDraw.Draw(im)
    for i in range(0, size[0], 16):  # stripes, so compression has real work to do
        d.rectangle((i, 0, i + 7, size[1]), fill=(i % 255, 120, 255 - i % 255))
    return im


@pytest.fixture
def jpg(tmp_path) -> Path:
    p = tmp_path / "photo.jpg"
    _draw((1600, 1000)).save(p, quality=95, exif=_gps_exif())
    return p


@pytest.fixture
def png(tmp_path) -> Path:
    p = tmp_path / "shot.png"
    _draw(mode="RGBA").save(p)
    return p


@pytest.fixture
def weird(tmp_path) -> Path:
    """Spaces, unicode, '#', '%20' and a leading dash all in one name."""
    p = tmp_path / "-tết #1 100%20 ảnh.png"
    _draw().save(p)
    return p


@pytest.fixture
def pdf(tmp_path, png) -> Path:
    fitz = eng.convert.fitz
    doc = fitz.open()
    for i in range(2):
        page = doc.new_page()
        page.insert_text((72, 72), f"Hello page {i + 1}", fontsize=24)
        page.insert_image(fitz.Rect(72, 150, 300, 300), filename=str(png))
    doc.set_metadata({"author": "Someone", "title": "Demo"})
    p = tmp_path / "Doc.pdf"
    doc.save(p)
    return p


@pytest.fixture
def pdf2(tmp_path) -> Path:
    doc = eng.convert.fitz.open()
    doc.new_page().insert_text((72, 72), "Second document", fontsize=20)
    p = tmp_path / "Other.pdf"
    doc.save(p)
    return p


@pytest.fixture
def docx(tmp_path) -> Path:
    import docx as pydocx
    d = pydocx.Document()
    d.add_heading("Title", 1)
    d.add_paragraph("Hello from a docx.")
    p = tmp_path / "Letter.docx"
    d.save(p)
    return p


@pytest.fixture
def txt(tmp_path) -> Path:
    p = tmp_path / "notes.md"
    p.write_text("# Notes\n\nSome *markdown* text.\n", encoding="utf-8")
    return p


@pytest.fixture(autouse=True)
def _private_state(tmp_path_factory, monkeypatch):
    """The created-folders ledger must never touch the real ~/.local/state."""
    monkeypatch.setenv("XDG_STATE_HOME", str(tmp_path_factory.mktemp("state")))


@pytest.fixture
def photos(tmp_path) -> Path:
    """A folder with nested images, a non-image, a hidden folder and a link."""
    d = tmp_path / "My Photos ảnh"
    (d / "trip" / "day 1").mkdir(parents=True)
    (d / ".thumbnails").mkdir()
    _draw((300, 200)).save(d / "a.png")
    _draw((300, 200)).save(d / "trip" / "b.jpg", quality=95)
    _draw((300, 200)).save(d / "trip" / "day 1" / "c #1.png")
    _draw((64, 64)).save(d / ".thumbnails" / "hidden.png")
    (d / "notes.txt").write_text("not an image\n")
    (d / "link.png").symlink_to(d / "a.png")
    return d
